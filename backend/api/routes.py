import csv
import io
import os
import threading
from datetime import datetime
from typing import List, Optional

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import FileResponse, JSONResponse, StreamingResponse
from sqlalchemy.orm import Session

from api import schemas
from database import database, models
from mailer.dispatcher import enviar_alertas_pendientes, enviar_correo, enviar_recordatorios
from scraper.anm_scraper import (BarridoEnCurso, barrido_en_curso, iniciar_barrido_en_segundo_plano,
                                 procesar_historico,
                                 REGIONALES_ANM, URL_BASE, URL_COACTIVO, descubrir_formulario)
from servicios.analizador import normalizar
from servicios.telegram import enviar_telegram, telegram_configurado
from servicios import plazos

router = APIRouter()


def _limpiar_regional(valor: Optional[str]) -> str:
    partes = [normalizar(p).strip() for p in (valor or "").replace(";", ",").split(",") if p.strip()]
    partes = [p for p in partes if p in REGIONALES_ANM or p == "all"]
    return ",".join(dict.fromkeys(partes)) or "bucaramanga"


def _limpiar_texto(v: Optional[str]) -> Optional[str]:
    v = (v or "").strip()
    return v or None


# ============================================================ CLIENTES (PLACAS)

@router.get("/clientes", response_model=List[schemas.ClienteResponse])
def obtener_clientes(db: Session = Depends(database.get_db)):
    return db.query(models.Cliente).order_by(models.Cliente.placa.asc()).all()


@router.post("/clientes", response_model=schemas.ClienteResponse)
def crear_cliente(cliente: schemas.ClienteCreate, db: Session = Depends(database.get_db)):
    placa_limpia = cliente.placa.strip().upper()
    if db.query(models.Cliente).filter(models.Cliente.placa == placa_limpia).first():
        raise HTTPException(status_code=400, detail=f"La placa {placa_limpia} ya está registrada.")
    nuevo = models.Cliente(
        placa=placa_limpia,
        nombre_empresa=cliente.nombre_empresa.strip() if cliente.nombre_empresa else "SIN NOMBRE",
        regional=_limpiar_regional(cliente.regional),
        activo=cliente.activo,
        email=_limpiar_texto(cliente.email),
        nit=_limpiar_texto(cliente.nit),
    )
    db.add(nuevo)
    db.commit()
    db.refresh(nuevo)
    return nuevo


@router.put("/clientes/{id_cliente}", response_model=schemas.ClienteResponse)
def actualizar_cliente(id_cliente: int, datos: schemas.ClienteUpdate, db: Session = Depends(database.get_db)):
    cliente = db.query(models.Cliente).filter(models.Cliente.id_cliente == id_cliente).first()
    if not cliente:
        raise HTTPException(status_code=404, detail="Cliente no encontrado.")

    if datos.placa is not None:
        nueva_placa = datos.placa.strip().upper()
        if nueva_placa != cliente.placa:
            if db.query(models.Cliente).filter(models.Cliente.placa == nueva_placa).first():
                raise HTTPException(status_code=400, detail="Otra placa ya tiene ese código.")
            cliente.placa = nueva_placa
    if datos.nombre_empresa is not None:
        cliente.nombre_empresa = datos.nombre_empresa.strip()
    if datos.regional is not None:
        cliente.regional = _limpiar_regional(datos.regional)
    if datos.activo is not None:
        cliente.activo = datos.activo
    if datos.email is not None:
        cliente.email = _limpiar_texto(datos.email)
    if datos.nit is not None:
        cliente.nit = _limpiar_texto(datos.nit)

    db.commit()
    db.refresh(cliente)
    return cliente


@router.delete("/clientes/{id_cliente}")
def eliminar_cliente(id_cliente: int, db: Session = Depends(database.get_db)):
    cliente = db.query(models.Cliente).filter(models.Cliente.id_cliente == id_cliente).first()
    if not cliente:
        raise HTTPException(status_code=404, detail="Cliente no encontrado.")
    db.query(models.Notificacion).filter(models.Notificacion.id_cliente == id_cliente).delete()
    db.delete(cliente)
    db.commit()
    return {"mensaje": f"Placa {cliente.placa} y su historial eliminados correctamente."}


COLUMNAS_IMPORTACION = {
    "placa": ["placa", "titulo", "titulo minero", "codigo"],
    "nombre_empresa": ["titular", "nombre", "empresa", "nombre empresa", "nombre_empresa", "razon social"],
    "regional": ["regional", "region", "punto de atencion", "par"],
    "email": ["email", "correo", "correo electronico", "e-mail"],
    "nit": ["nit", "cc", "nit/cc", "cedula", "documento"],
}


def _leer_archivo_importacion(nombre: str, contenido: bytes):
    if nombre.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        hoja = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True).active
        filas = [[("" if c is None else str(c)).strip() for c in fila] for fila in hoja.iter_rows(values_only=True)]
    elif nombre.lower().endswith(".csv"):
        texto = contenido.decode("utf-8-sig", errors="replace")
        dialecto = csv.Sniffer().sniff(texto[:2000], delimiters=",;\t") if texto.strip() else csv.excel
        filas = [[c.strip() for c in fila] for fila in csv.reader(io.StringIO(texto), dialecto)]
    else:
        raise HTTPException(status_code=400, detail="Formato no soportado. Use .xlsx o .csv")
    filas = [f for f in filas if any(f)]
    if not filas:
        raise HTTPException(status_code=400, detail="El archivo está vacío.")

    encabezados = [normalizar(h).strip() for h in filas[0]]
    indices = {}
    for campo, alias in COLUMNAS_IMPORTACION.items():
        for i, h in enumerate(encabezados):
            if h in alias:
                indices[campo] = i
                break
    if "placa" not in indices:
        raise HTTPException(status_code=400, detail="No se encontró la columna 'placa' en la primera fila.")
    return [{campo: (fila[i] if i < len(fila) else "") for campo, i in indices.items()} for fila in filas[1:]]


@router.post("/clientes/importar")
async def importar_clientes(archivo: UploadFile = File(...), db: Session = Depends(database.get_db)):
    """Crea o actualiza placas desde Excel/CSV con columnas: placa, titular, regional, email, nit."""
    registros = _leer_archivo_importacion(archivo.filename or "", await archivo.read())
    creados, actualizados, omitidos = 0, 0, []
    for i, r in enumerate(registros, start=2):
        placa = (r.get("placa") or "").strip().upper()
        if not placa:
            omitidos.append(f"Fila {i}: sin placa")
            continue
        existente = db.query(models.Cliente).filter(models.Cliente.placa == placa).first()
        if existente:
            if r.get("nombre_empresa"):
                existente.nombre_empresa = r["nombre_empresa"].strip().upper()
            if r.get("regional"):
                existente.regional = _limpiar_regional(r["regional"])
            if r.get("email"):
                existente.email = r["email"].strip()
            if r.get("nit"):
                existente.nit = r["nit"].strip()
            actualizados += 1
        else:
            db.add(models.Cliente(
                placa=placa,
                nombre_empresa=(r.get("nombre_empresa") or "SIN NOMBRE").strip().upper(),
                regional=_limpiar_regional(r.get("regional")),
                activo=True,
                email=_limpiar_texto(r.get("email")),
                nit=_limpiar_texto(r.get("nit")),
            ))
            creados += 1
    db.commit()
    return {"creados": creados, "actualizados": actualizados, "omitidos": omitidos}


@router.get("/clientes/plantilla")
def plantilla_importacion():
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Placas"
    ws.append(["placa", "titular", "regional", "email", "nit"])
    ws.append(["HAN-111", "EMPRESA EJEMPLO SAS", "bucaramanga", "titular@ejemplo.com", "900123456"])
    ws.append(["AB1-234", "OTRO TITULAR", "bogota,cartagena", "", ""])
    for col, ancho in zip("ABCDE", [14, 34, 22, 30, 14]):
        ws.column_dimensions[col].width = ancho
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": 'attachment; filename="plantilla_placas_anm.xlsx"'})


# ============================================================ NOTIFICACIONES

@router.get("/notificaciones", response_model=List[schemas.NotificacionResponse])
def obtener_notificaciones(db: Session = Depends(database.get_db)):
    return db.query(models.Notificacion).order_by(models.Notificacion.id_notificacion.desc()).all()


@router.get("/vencimientos", response_model=List[schemas.NotificacionResponse])
def obtener_vencimientos(incluir_vencidas_dias: int = 10, db: Session = Depends(database.get_db)):
    """Notificaciones no respondidas con fecha límite próxima (y las vencidas recientes)."""
    from datetime import timedelta
    desde = plazos.hoy_co() - timedelta(days=incluir_vencidas_dias)
    return (db.query(models.Notificacion)
            .filter(models.Notificacion.fecha_limite != None, models.Notificacion.fecha_limite >= desde,
                    models.Notificacion.estado != "respondida")
            .order_by(models.Notificacion.fecha_limite.asc()).all())


@router.put("/notificaciones/{id_notificacion}", response_model=schemas.NotificacionResponse)
def gestionar_notificacion(id_notificacion: int, datos: schemas.NotificacionGestion,
                           db: Session = Depends(database.get_db)):
    notif = db.query(models.Notificacion).filter(models.Notificacion.id_notificacion == id_notificacion).first()
    if not notif:
        raise HTTPException(status_code=404, detail="Notificación no encontrada.")
    if datos.estado is not None:
        if datos.estado not in models.ESTADOS_GESTION:
            raise HTTPException(status_code=400, detail=f"Estado inválido. Use: {', '.join(models.ESTADOS_GESTION)}")
        notif.estado = datos.estado
    if datos.notas is not None:
        notif.notas = datos.notas.strip() or None
    notif.fecha_gestion = datetime.utcnow()
    db.commit()
    db.refresh(notif)
    return notif


@router.get("/notificaciones/{id_notificacion}/pdf")
def descargar_copia_pdf(id_notificacion: int, db: Session = Depends(database.get_db)):
    notif = db.query(models.Notificacion).filter(models.Notificacion.id_notificacion == id_notificacion).first()
    if not notif or not notif.tiene_archivo:
        raise HTTPException(status_code=404, detail="No hay copia archivada de este documento.")
    return FileResponse(notif.archivo_local, media_type="application/pdf",
                        filename=os.path.basename(notif.archivo_local), content_disposition_type="inline")


@router.post("/notificaciones/{id_notificacion}/resumir", response_model=schemas.NotificacionResponse)
def resumir_notificacion(id_notificacion: int, db: Session = Depends(database.get_db)):
    """Genera (o regenera) el resumen con IA de una notificación."""
    from servicios.analizador import generar_resumen_ia
    from servicios.procesamiento import texto_de_notificacion
    notif = db.query(models.Notificacion).filter(models.Notificacion.id_notificacion == id_notificacion).first()
    if not notif:
        raise HTTPException(status_code=404, detail="Notificación no encontrada.")
    try:
        texto = texto_de_notificacion(notif)
    except Exception as e:
        raise HTTPException(status_code=502, detail=str(e))
    resumen, error = generar_resumen_ia(texto)
    if error:
        db.commit()  # conserva la copia archivada si se acaba de descargar
        raise HTTPException(status_code=502, detail=error)
    notif.resumen = resumen
    db.commit()
    db.refresh(notif)
    return notif


@router.delete("/notificaciones/{id_notificacion}")
def eliminar_notificacion(id_notificacion: int, db: Session = Depends(database.get_db)):
    notif = db.query(models.Notificacion).filter(models.Notificacion.id_notificacion == id_notificacion).first()
    if not notif:
        raise HTTPException(status_code=404, detail="Notificación no encontrada.")
    db.delete(notif)
    db.commit()
    return {"mensaje": "Notificación eliminada correctamente."}


@router.post("/notificaciones/marcar-respondidas")
def marcar_antiguas_respondidas(antes_de: str, db: Session = Depends(database.get_db)):
    """Marca como respondidas (y ya notificadas) las notificaciones publicadas antes de una fecha."""
    fecha = plazos.parsear_fecha(antes_de)
    if not fecha:
        raise HTTPException(status_code=400, detail="Fecha inválida. Use el formato AAAA-MM-DD.")
    antiguas = [n for n in db.query(models.Notificacion).filter(models.Notificacion.estado != "respondida").all()
                if (plazos.parsear_fecha(n.fecha_aviso) or fecha) < fecha]
    for n in antiguas:
        n.estado = "respondida"
        n.notificacion_enviada = True
        n.fecha_gestion = datetime.utcnow()
        if not n.notas:
            n.notas = f"Marcada como resuelta en bloque (anterior a {fecha.isoformat()})"
    db.commit()
    return {"marcadas": len(antiguas)}


@router.post("/notificaciones/enviar-pendientes")
def reenviar_pendientes(db: Session = Depends(database.get_db)):
    resultado = enviar_alertas_pendientes(db)
    if resultado.get("error"):
        raise HTTPException(status_code=500, detail=resultado["error"])
    return resultado


# ============================================================ BARRIDOS

@router.post("/scraper/ejecutar", status_code=202)
def disparar_scraper(id_cliente: Optional[int] = None, db: Session = Depends(database.get_db)):
    """
    Lanza un barrido en segundo plano y responde de inmediato (202 Accepted) con el id
    de la ejecución. El avance se consulta en GET /scraper/ejecuciones/{id}.
    """
    if id_cliente is not None and db.get(models.Cliente, id_cliente) is None:
        raise HTTPException(status_code=404, detail="Título minero no encontrado.")
    try:
        id_ejecucion = iniciar_barrido_en_segundo_plano(id_cliente_especifico=id_cliente, origen="manual")
    except BarridoEnCurso as e:
        # Se devuelve el id del barrido que ya corre para que el panel pueda seguirlo en vez de fallar
        actual = barrido_en_curso()
        return JSONResponse(status_code=409, content={
            "detail": str(e), "id_ejecucion": actual.id_ejecucion if actual else None})
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"No se pudo iniciar el barrido: {e}")
    return {"id_ejecucion": id_ejecucion, "estado": "en_curso"}


@router.get("/scraper/estado", response_model=Optional[schemas.EjecucionResponse])
def estado_barrido_actual():
    """El barrido que está corriendo ahora (manual o del cron), o null si no hay ninguno."""
    return barrido_en_curso()


@router.get("/scraper/ejecuciones/{id_ejecucion}", response_model=schemas.EjecucionResponse)
def estado_de_ejecucion(id_ejecucion: int, db: Session = Depends(database.get_db)):
    """Avance de un barrido. Ruta abierta al usuario del panel (a diferencia de /ejecuciones)."""
    ejec = db.get(models.Ejecucion, id_ejecucion)
    if ejec is None:
        raise HTTPException(status_code=404, detail="Ejecución no encontrada.")
    return ejec


@router.get("/ejecuciones", response_model=List[schemas.EjecucionResponse])
def obtener_ejecuciones(limite: int = 40, db: Session = Depends(database.get_db)):
    return db.query(models.Ejecucion).order_by(models.Ejecucion.inicio.desc()).limit(limite).all()


_estado_historico = {"en_curso": False, "resultado": None, "inicio": None}


@router.post("/mantenimiento/procesar-historico")
def iniciar_procesamiento_historico(usar_ia: bool = False):
    """Archiva, analiza y calcula plazos de las notificaciones antiguas (en segundo plano)."""
    if _estado_historico["en_curso"]:
        return {"mensaje": "Ya se está procesando el histórico.", **_estado_historico}

    def tarea():
        _estado_historico.update(en_curso=True, resultado=None, inicio=datetime.utcnow().isoformat())
        try:
            _estado_historico["resultado"] = procesar_historico(usar_ia=usar_ia)
        except Exception as e:
            _estado_historico["resultado"] = {"error": str(e)}
        finally:
            _estado_historico["en_curso"] = False

    threading.Thread(target=tarea, daemon=True).start()
    return {"mensaje": "Procesamiento del histórico iniciado en segundo plano."}


@router.get("/mantenimiento/procesar-historico")
def estado_procesamiento_historico(db: Session = Depends(database.get_db)):
    sin_procesar = db.query(models.Notificacion).filter(
        (models.Notificacion.procesada == False) | (models.Notificacion.procesada == None)).count()
    return {**_estado_historico, "sin_procesar": sin_procesar}


@router.post("/recordatorios/ejecutar")
def ejecutar_recordatorios(db: Session = Depends(database.get_db)):
    return enviar_recordatorios(db)


# ============================================================ DIAGNÓSTICO

@router.get("/correo/diagnostico")
def diagnostico_correo(db: Session = Depends(database.get_db)):
    pendientes = db.query(models.Notificacion).filter(models.Notificacion.notificacion_enviada == False).count()
    return {
        "modo_envio": "brevo_api" if os.getenv("BREVO_API_KEY") else ("gmail_smtp" if os.getenv("EMAIL_PASSWORD") else "sin_configurar"),
        "BREVO_API_KEY": bool(os.getenv("BREVO_API_KEY")),
        "EMAIL_SENDER": bool(os.getenv("EMAIL_SENDER")),
        "EMAIL_RECEIVER": bool(os.getenv("EMAIL_RECEIVER")),
        "EMAIL_PASSWORD": bool(os.getenv("EMAIL_PASSWORD")),
        "telegram": telegram_configurado(),
        "usuario_panel": bool(os.getenv("PANEL_USER") and os.getenv("PANEL_PASSWORD")),
        "resumen_ia_gemini": bool(os.getenv("GEMINI_API_KEY")),
        "base_de_datos": "postgres" if os.getenv("DATABASE_URL") else ("volume" if os.getenv("DB_PATH") else "repositorio (se reinicia en cada deploy)"),
        "notificaciones_pendientes": pendientes,
    }


@router.post("/alertas/prueba")
def probar_alertas():
    """Envía un correo y un mensaje de Telegram de prueba."""
    resultado = {}
    try:
        enviar_correo(os.getenv("EMAIL_RECEIVER"), "✅ Prueba del sistema de monitoreo ANM",
                      "<p>Si recibe este correo, el envío de notificaciones está funcionando.</p>")
        resultado["correo"] = "enviado"
    except Exception as e:
        resultado["correo"] = f"error: {e}"
    if telegram_configurado():
        resultado["telegram"] = "enviado" if enviar_telegram("✅ Prueba del sistema de monitoreo ANM") else "error (ver logs)"
    else:
        resultado["telegram"] = "no configurado"
    return resultado


@router.get("/salud/fuentes")
def revisar_fuentes():
    """Comprueba en vivo que las páginas de la ANM conserven la estructura que el scraper espera."""
    import requests
    sesion = requests.Session()
    sesion.headers["User-Agent"] = "Mozilla/5.0 (Windows NT 10.0; Win64; x64) Chrome/122.0 Safari/537.36"
    salida = {}
    try:
        r = sesion.get(URL_BASE, timeout=20)
        salida["avisos"] = {"http": r.status_code, "filtro_placa_encontrado": "field_pla_tit_min" in r.text,
                            "tabla_encontrada": "<table" in r.text}
    except Exception as e:
        salida["avisos"] = {"error": str(e)}
    try:
        r = sesion.get(URL_COACTIVO, timeout=20)
        campos = descubrir_formulario(r.text)
        salida["cobro_coactivo"] = {"http": r.status_code, "campos_formulario": campos,
                                    "campo_palabra_clave": any("palabra" in k or "clave" in k for k in campos)}
    except Exception as e:
        salida["cobro_coactivo"] = {"error": str(e)}
    return salida
